#!/usr/bin/env python3
"""Black Minimal Deck — render a deck HTML file to a 16:9 PPTX.

Pipeline: headless Chromium screenshots each slide at 3840x2160 (viewport
1920x1080 @ deviceScaleFactor=2), then python-pptx assembles them as
full-bleed images on 16:9 pages. Text is preserved as high-res bitmap:
zero distortion, zero dropped glyphs — at the cost of non-editable text.

Usage:
    python scripts/render_pptx.py templates/deck.html my-deck.pptx
    python scripts/render_pptx.py deck.html out.pptx --pages 1-5,8 --scale 2

Deps:
    pip install playwright python-pptx
    playwright install chromium      # not needed if system Chrome exists
"""

from __future__ import annotations

import argparse
import hashlib
import sys
import tempfile
from pathlib import Path


def parse_pages(spec: str | None, total: int) -> list[int]:
    """'1-3,5' -> [1,2,3,5]; None -> all pages. 1-based."""
    if not spec:
        return list(range(1, total + 1))
    pages: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            pages.extend(range(int(a), int(b) + 1))
        elif part:
            pages.append(int(part))
    out = []
    for n in pages:
        if 1 <= n <= total and n not in out:
            out.append(n)
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description="Render a Black Minimal Deck HTML to PPTX")
    ap.add_argument("html", help="deck HTML file (templates/deck.html)")
    ap.add_argument("pptx", help="output .pptx path")
    ap.add_argument("--pages", default=None, help="pages to export, e.g. 1-5,8 (default: all)")
    ap.add_argument("--scale", type=int, default=2, help="device scale factor (default 2 => 3840x2160)")
    ap.add_argument("--wait-ms", type=int, default=1400, help="per-slide settle time for fonts/anims")
    args = ap.parse_args()

    html = Path(args.html).resolve()
    if not html.exists():
        sys.exit(f"[render_pptx] file not found: {html}")

    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        sys.exit("[render_pptx] missing deps:\n  pip install playwright python-pptx\n  playwright install chromium")

    from pptx import Presentation
    from pptx.util import Inches

    url = html.as_uri()

    with tempfile.TemporaryDirectory(prefix="bmd_render_") as td:
        shots: list[Path] = []

        with sync_playwright() as p:
            # Prefer system Chrome (no browser download); fall back to bundled Chromium.
            try:
                browser = p.chromium.launch(channel="chrome", args=["--no-sandbox", "--disable-gpu"])
            except Exception:
                browser = p.chromium.launch(args=["--no-sandbox", "--disable-gpu"])

            page = browser.new_page(
                viewport={"width": 1920, "height": 1080},
                device_scale_factor=args.scale,
            )

            page.goto(url, wait_until="load")
            try:
                page.wait_for_load_state("networkidle", timeout=8000)
            except Exception:
                pass  # offline mode: Google Fonts unreachable, system fallback kicks in
            page.wait_for_timeout(args.wait_ms)
            total = page.evaluate("document.querySelectorAll('.slide').length")
            if not total:
                browser.close()
                sys.exit("[render_pptx] no .slide sections found in the HTML")

            pages = parse_pages(args.pages, total)
            if not pages:
                browser.close()
                sys.exit(f"[render_pptx] --pages {args.pages} matches no slide (valid range: 1-{total})")
            print(f"[render_pptx] {total} slides found, exporting pages: {pages}")

            # Hide on-screen nav controls so they never appear in the export.
            page.add_style_tag(content=".nav-ctrl{display:none!important}")

            # Activate each slide in place. Do NOT rely on "#slide-N" navigation:
            # a fragment change is same-document, so the deck JS never re-runs and
            # every capture silently renders the same page.
            activate = """(idx) => {
              const ss = [...document.querySelectorAll('.slide')];
              ss.forEach(s => s.classList.remove('active'));
              const el = document.getElementById('slide-' + idx);
              if (!el) return false;
              el.classList.add('active');
              return true;
            }"""

            for n in pages:
                if not page.evaluate(activate, n):
                    print(f"[render_pptx] WARNING: no slide with id 'slide-{n}', skipped")
                    continue
                page.wait_for_timeout(args.wait_ms)
                shot = Path(td) / f"slide-{n:02d}.png"
                page.screenshot(path=str(shot), clip={"x": 0, "y": 0, "width": 1920, "height": 1080})
                shots.append(shot)
                print(f"[render_pptx] captured page {n}/{total}")

            # Guard against the "same page exported N times" failure mode.
            digests = {hashlib.md5(s.read_bytes()).hexdigest() for s in shots}
            if len(shots) > 1 and len(digests) == 1:
                print("[render_pptx] WARNING: every captured page is identical - slide activation likely failed")

            browser.close()

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]
        for shot in shots:
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(shot), 0, 0, width=prs.slide_width, height=prs.slide_height)

        out = Path(args.pptx)
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
        print(f"[render_pptx] saved {len(shots)} pages -> {out}")


if __name__ == "__main__":
    main()
